import base64
import unicodedata
from io import BytesIO
from unittest.mock import MagicMock, patch
from zipfile import ZipFile

from lxml import etree

from odoo import fields
from odoo.addons.ranvals_document_studio.models.rds_template import FONT_CSS
from odoo.addons.ranvals_document_studio.tools.docx_editable_renderer import (
    EDITABLE_DOCX_FINGERPRINT,
    MAX_ADDRESS_LINES,
    MAX_COLUMNS,
    MAX_LINES,
    MAX_LOGO_DIMENSION,
    MAX_METADATA_ITEMS,
    MAX_NOTES,
    MAX_TEXT_CHARS,
    MAX_TOTALS,
    EditableDocxError,
    _font,
    _resolve_docx_locale,
    render_editable_docx,
)
from odoo.addons.ranvals_document_studio.tools.docx_renderer import (
    DOCX_RENDERER_FINGERPRINT,
    MAX_DOCX_BYTES,
    render_docx,
)
from odoo.addons.ranvals_document_studio.wizard.rds_export_wizard import (
    DOCX_MIMETYPE,
    MAX_EXPORT_BYTES,
    RdsExportWizard,
)
from odoo.exceptions import AccessError, UserError, ValidationError
from odoo.tests.common import TransactionCase, new_test_user, tagged
from PIL import Image


@tagged("post_install", "-at_install")
class TestRdsExportWizard(TransactionCase):

    def setUp(self):
        super().setUp()
        self.wizard = self.env["rds.export.wizard"].create(
            {
                "res_model": "res.partner",
                "res_ids_json": "[%s]" % self.env.user.partner_id.id,
                "template_id": self.env.ref(
                    "ranvals_document_studio.template_technology_blue"
                ).id,
                "output_format": "docx",
            }
        )

    @staticmethod
    def _editable_context():
        return {
            "title": "Teklif",
            "number": "S-1",
            "theme": {
                "primary": "#17345F",
                "secondary": "#F4F7FA",
                "accent": "#D33A35",
                "text": "#27313A",
            },
            "company_info": {
                "name": "ŞİRKET-GÖRÜNÜR",
                "lines": ["ŞİRKET-ADRESİ"],
                "email": "company@example.com",
                "vat": "ŞİRKET-VKN-123",
            },
            "partner_info": {
                "name": "PARTNER-GÖRÜNÜR",
                "lines": ["PARTNER-ADRESİ"],
                "email": "partner@example.com",
            },
            "labels": {
                "company_info": "Şirket Bilgileri",
                "partner_info": "Müşteri Bilgileri",
                "document_info": "Belge Bilgileri",
                "notes": "Notlar",
                "bank_info": "Banka Bilgileri",
                "phone": "Telefon",
                "email": "E-posta",
                "website": "Web Sitesi",
                "tax_no": "Vergi No",
            },
            "tax_label": "Vergi No",
            "metadata": [
                {"label": "Meta", "value": "META-GÖRÜNÜR", "align": "center"}
            ],
            "columns": [
                {"label": "Açıklama", "align": "left", "width": 60},
                {"label": "Tutar", "align": "right", "width": 40},
            ],
            "lines": [
                {
                    "values": [
                        {"text": "Hizmet", "align": "left"},
                        {"text": "100,00 TRY", "align": "right"},
                    ]
                }
            ],
            "totals": [
                {"label": "Toplam", "value": "100,00 TRY", "is_total": True}
            ],
            "notes": ["Ödeme 30 gün"],
            "bank": {"iban": "TR00 0000 0000", "swift": "ABCDEF12"},
            "bank_labels": {"iban": "IBAN", "swift": "SWIFT"},
            "footer_text": "Teşekkürler",
        }

    def test_export_uses_popup_free_client_download(self):
        content = b"PK\x03\x04test-docx"
        mimetype = (
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        )

        with patch.object(
            RdsExportWizard,
            "_build_output",
            return_value=(
                "test-document.docx",
                content,
                mimetype,
                [
                    (
                        self.env.user.partner_id,
                        [("test-document.docx", content, mimetype)],
                    )
                ],
            ),
        ):
            action = self.wizard.action_export()

        self.assertEqual(action["type"], "ir.actions.client")
        self.assertEqual(
            action["tag"], "ranvals_document_studio.download_export"
        )
        self.assertEqual(
            action["params"]["url"],
            "/ranvals_document_studio/export_wizard/%s" % self.wizard.id,
        )
        self.assertEqual(base64.b64decode(self.wizard.file_data), content)
        self.assertEqual(self.wizard.file_name, "test-document.docx")
        self.assertEqual(self.wizard.file_mimetype, mimetype)
        export_log = self.env["rds.export.log"].search(
            [("file_name", "=", "test-document.docx")], limit=1
        )
        self.assertTrue(export_log.has_stored_file)
        self.assertFalse(export_log.attached_to_record)
        self.assertEqual(export_log.file_mimetype, mimetype)
        self.assertEqual(export_log.file_size, len(content))
        self.assertEqual(base64.b64decode(export_log.file_data), content)
        export_log.invalidate_recordset(["file_data"])
        self.assertEqual(base64.b64decode(export_log.file_data), content)
        stored_attachment = self.env["ir.attachment"].search(
            [
                ("res_model", "=", "rds.export.log"),
                ("res_field", "=", "file_data"),
                ("res_id", "=", export_log.id),
            ]
        )
        self.assertEqual(len(stored_attachment), 1)

    def test_editable_word_proofing_locale_uses_supported_language_prefix(self):
        self.assertEqual(_resolve_docx_locale("de_AT"), "de-DE")
        self.assertEqual(_resolve_docx_locale("es_MX"), "es-ES")
        self.assertEqual(_resolve_docx_locale("fr_CA"), "fr-FR")
        self.assertEqual(_resolve_docx_locale("ar_AE"), "ar-SA")
        self.assertEqual(_resolve_docx_locale("ja_JP"), "en-US")

    def test_pdf_preview_uses_popup_free_dialog(self):
        self.wizard.output_format = "pdf"
        content = b"%PDF-test"
        mimetype = "application/pdf"
        with patch.object(
            RdsExportWizard,
            "_build_output",
            return_value=(
                "preview.pdf",
                content,
                mimetype,
                [
                    (
                        self.env.user.partner_id,
                        [("preview.pdf", content, mimetype)],
                    )
                ],
            ),
        ):
            action = self.wizard._save_result(download=False)

        self.assertEqual(action["type"], "ir.actions.client")
        self.assertEqual(action["tag"], "ranvals_document_studio.preview_export")
        self.assertEqual(
            action["params"]["url"],
            "/ranvals_document_studio/export_wizard/%s/preview" % self.wizard.id,
        )
        self.assertEqual(action["params"]["title"], "preview.pdf Preview")
        self.assertFalse(
            self.env["rds.export.log"].search(
                [("file_name", "=", "preview.pdf")]
            )
        )

    def test_docx_renderer_produces_valid_full_page_a4_ooxml(self):
        page_stream = BytesIO()
        Image.new("RGB", (2480, 3508), (244, 248, 255)).save(
            page_stream, format="PNG", dpi=(300, 300)
        )

        with patch(
            "odoo.addons.ranvals_document_studio.tools.docx_renderer.render_png_pages",
            return_value=[page_stream.getvalue()],
        ) as render_pages:
            content = render_docx(b"%PDF-test")

        self.assertTrue(content.startswith(b"PK\x03\x04"))
        with ZipFile(BytesIO(content)) as archive:
            self.assertIsNone(archive.testzip())
            document_xml = archive.read("word/document.xml")
            settings_xml = archive.read("word/settings.xml")
            core_xml = archive.read("docProps/core.xml")
            media_name = next(
                name for name in archive.namelist() if name.startswith("word/media/")
            )
            embedded_page = Image.open(BytesIO(archive.read(media_name)))
        root = etree.fromstring(
            document_xml,
            parser=etree.XMLParser(resolve_entities=False, no_network=True),
        )
        word_namespace = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
        drawing_namespace = (
            "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
        )
        page_size = root.find(".//{%s}pgSz" % word_namespace)
        page_margins = root.find(".//{%s}pgMar" % word_namespace)
        page_anchor = root.find(".//{%s}anchor" % drawing_namespace)
        page_extent = root.find(".//{%s}extent" % drawing_namespace)
        self.assertIsNotNone(page_size)
        self.assertEqual(page_size.get("{%s}w" % word_namespace), "11906")
        self.assertEqual(page_size.get("{%s}h" % word_namespace), "16838")
        self.assertTrue(all(
            page_margins.get("{%s}%s" % (word_namespace, edge)) == "0"
            for edge in ("top", "right", "bottom", "left")
        ))
        self.assertIsNotNone(page_anchor)
        self.assertNotIn(b"<wp:inline", document_xml)
        self.assertNotIn(b"<w:tbl", document_xml)
        self.assertEqual(page_extent.get("cx"), "7560000")
        self.assertEqual(page_extent.get("cy"), "10692000")
        self.assertIn(b"doNotAutoCompressPictures", settings_xml)
        self.assertIn(b"defaultImageDpi", settings_xml)
        self.assertIn(DOCX_RENDERER_FINGERPRINT.encode(), core_xml)
        self.assertEqual(embedded_page.size, (2480, 3508))
        render_pages.assert_called_once_with(
            b"%PDF-test",
            dpi=300,
            max_pages=50,
            max_output_bytes=MAX_DOCX_BYTES,
        )

    def test_docx_export_uses_qweb_pdf_once_and_skips_legacy_context(self):
        record = self.env.user.partner_id
        with (
            patch.object(
                RdsExportWizard,
                "_render_pdf",
                return_value=b"%PDF-same-source",
            ) as render_pdf,
            patch.object(RdsExportWizard, "_record_context") as record_context,
            patch(
                "odoo.addons.ranvals_document_studio.wizard.rds_export_wizard.render_docx",
                return_value=b"PK\x03\x04docx",
            ) as render_word,
            patch(
                "odoo.addons.ranvals_document_studio.wizard.rds_localization.localize_docx",
                side_effect=lambda content, _language: content,
            ),
        ):
            outputs = self.wizard._export_record(record, "tr_TR")

        render_pdf.assert_called_once_with(record, "tr_TR")
        record_context.assert_not_called()
        render_word.assert_called_once_with(b"%PDF-same-source", max_pages=50)
        self.assertEqual(len(outputs), 1)
        self.assertTrue(outputs[0][0].endswith(".docx"))
        self.assertEqual(
            outputs[0][2],
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )

    def test_editable_docx_contains_native_text_tables_and_no_page_raster(self):
        context = {
            "title": "Satış Teklifi",
            "number": "S00042",
            "direction": "ltr",
            "lang_code": "tr",
            "theme": {
                "primary": "#17345F",
                "secondary": "#F4F7FA",
                "accent": "#D33A35",
                "text": "#27313A",
                "heading_font": "Arial, sans-serif",
                "body_font": "Arial, sans-serif",
            },
            "company_info": {
                "name": "Ranvals AŞ",
                "lines": ["Teknoloji Cad. 1", "İstanbul"],
                "phone": "+90 212 000 00 00",
            },
            "partner_info": {
                "name": "Örnek Müşteri",
                "lines": ["Müşteri Sok. 2", "Ankara"],
                "email": "musteri@example.com",
                "vat": "1234567890",
            },
            "labels": {
                "company_info": "Şirket Bilgileri",
                "partner_info": "Müşteri Bilgileri",
                "document_info": "Belge Bilgileri",
                "notes": "Notlar ve Ödeme Koşulları",
                "bank_info": "Banka Bilgileri",
                "phone": "Telefon",
                "email": "E-posta",
                "website": "Web Sitesi",
                "tax_no": "Vergi No",
            },
            "tax_label": "Vergi No",
            "metadata": [
                {"label": "Tarih", "value": "24.09.2026"},
                {"label": "Satış Temsilcisi", "value": "Demo Kullanıcı"},
            ],
            "columns": [
                {"label": "Açıklama", "align": "left", "width": 60},
                {"label": "Tutar", "align": "right", "width": 40},
            ],
            "lines": [
                {"is_section": True, "description": "Hizmetler"},
                {
                    "values": [
                        {"text": "Danışmanlık"},
                        {"text": "1.000,00 TL", "align": "right", "bold": True},
                    ]
                },
                {"is_note": True, "description": "Uzaktan teslim edilecektir."},
            ],
            "totals": [
                {"label": "Ara Toplam", "value": "1.000,00 TL"},
                {"label": "Genel Toplam", "value": "1.200,00 TL", "is_total": True},
            ],
            "notes": ["Ödeme vadesi 30 gündür."],
            "bank": {"bank_name": "Örnek Banka", "iban": "TR00 0000"},
            "bank_labels": {"bank": "Banka", "iban": "IBAN"},
            "footer_text": "Teşekkür ederiz.",
        }

        content = render_editable_docx(context, "tr_TR")

        self.assertTrue(content.startswith(b"PK\x03\x04"))
        with ZipFile(BytesIO(content)) as archive:
            self.assertIsNone(archive.testzip())
            document_xml = archive.read("word/document.xml")
            core_xml = archive.read("docProps/core.xml")
            media = [name for name in archive.namelist() if name.startswith("word/media/")]
        self.assertIn(b"<w:tbl", document_xml)
        self.assertIn("Danışmanlık".encode(), document_xml)
        self.assertIn("Ödeme vadesi 30 gündür.".encode(), document_xml)
        self.assertNotIn(b"<wp:anchor", document_xml)
        self.assertEqual(media, [])
        self.assertIn(EDITABLE_DOCX_FINGERPRINT.encode(), core_xml)

        # python-docx can reopen the result and sees ordinary editable cells,
        # rather than a full-page screenshot masquerading as a Word file.
        from docx import Document

        document = Document(BytesIO(content))
        table_text = "\n".join(
            cell.text
            for table in document.tables
            for row in table.rows
            for cell in row.cells
        )
        self.assertIn("Danışmanlık", table_text)
        self.assertIn("1.200,00 TL", table_text)
        self.assertGreaterEqual(len(document.tables), 4)
        self.assertEqual(document.core_properties.author, "Ranvals Software")
        self.assertEqual(document.core_properties.last_modified_by, "DocuCraft")
        self.assertEqual(document.core_properties.title, "Satış Teklifi S00042")
        company_cells = [
            cell
            for table in document.tables
            for row in table.rows
            for cell in row.cells
            if cell.paragraphs
            and cell.paragraphs[0].text == "Şirket Bilgileri"
        ]
        self.assertEqual(len(company_cells), 1)
        self.assertEqual(company_cells[0].paragraphs[1].text, "Ranvals AŞ")
        self.assertIn("Telefon: +90 212 000 00 00", table_text)
        self.assertIn("E-posta: musteri@example.com", table_text)

    def test_editable_docx_uses_document_language_for_contact_labels(self):
        context = self._editable_context()
        context["labels"].update({
            "phone": "Phone",
            "email": "Email",
            "website": "Website",
            "tax_no": "Tax ID",
        })
        context["tax_label"] = "Tax ID"

        from docx import Document

        document = Document(BytesIO(render_editable_docx(context, "en_US")))
        text = "\n".join(
            cell.text
            for table in document.tables
            for row in table.rows
            for cell in row.cells
        )
        self.assertIn("Email: company@example.com", text)
        self.assertIn("Email: partner@example.com", text)
        self.assertIn("Tax ID: ŞİRKET-VKN-123", text)
        self.assertNotIn("E-posta:", text)
        self.assertNotIn("Vergi No:", text)

    def test_editable_docx_supports_every_template_font_and_script_slot(self):
        expected_word_fonts = {
            "serif": "Georgia",
            "sans": "Arial",
            "technical": "Trebuchet MS",
            "editorial": "Georgia",
            "lato": "Lato",
            "roboto": "Roboto",
            "open_sans": "Open Sans",
            "montserrat": "Montserrat",
            "raleway": "Raleway",
            "oswald": "Oswald",
            "tajawal": "Tajawal",
            "fira_mono": "Fira Mono",
            "humanist": "Calibri",
            "helvetica": "Arial",
            "verdana": "Verdana",
            "tahoma": "Tahoma",
            "lucida": "Lucida Sans",
            "times": "Times New Roman",
            "garamond": "Garamond",
            "palatino": "Palatino Linotype",
            "cambria": "Cambria",
            "bookman": "Bookman Old Style",
            "monospace": "Courier New",
        }
        self.assertEqual(set(expected_word_fonts), set(FONT_CSS))
        for font_key, word_font in expected_word_fonts.items():
            with self.subTest(font_key=font_key):
                self.assertEqual(_font(font_key), word_font)
                self.assertEqual(_font(FONT_CSS[font_key]), word_font)

        context = self._editable_context()
        context["theme"].update({
            "heading_font": FONT_CSS["montserrat"],
            "body_font": FONT_CSS["open_sans"],
        })
        content = render_editable_docx(context, "tr_TR")
        with ZipFile(BytesIO(content)) as archive:
            xml_roots = [
                etree.fromstring(archive.read(part_name))
                for part_name in (
                    "word/document.xml",
                    "word/styles.xml",
                    "word/footer1.xml",
                )
            ]

        namespace = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
        attribute = "{%s}%%s" % namespace["w"]
        for word_font in ("Montserrat", "Open Sans"):
            font_nodes = [
                node
                for root in xml_roots
                for node in root.xpath(
                    "//w:rFonts[@w:ascii=$font]",
                    namespaces=namespace,
                    font=word_font,
                )
            ]
            self.assertTrue(font_nodes, word_font)
            for node in font_nodes:
                with self.subTest(word_font=word_font, xml=etree.tostring(node)):
                    self.assertEqual(
                        [node.get(attribute % slot) for slot in ("ascii", "hAnsi", "eastAsia", "cs")],
                        [word_font] * 4,
                    )

    def test_editable_docx_rejects_every_bounded_list_overflow(self):
        cases = {
            "lines": lambda context: context.update(
                lines=[{}] * (MAX_LINES + 1)
            ),
            "columns": lambda context: context.update(
                columns=[{}] * (MAX_COLUMNS + 1)
            ),
            "metadata": lambda context: context.update(
                metadata=[{}] * (MAX_METADATA_ITEMS + 1)
            ),
            "notes": lambda context: context.update(
                notes=[""] * (MAX_NOTES + 1)
            ),
            "totals": lambda context: context.update(
                totals=[{}] * (MAX_TOTALS + 1)
            ),
            "address": lambda context: context["company_info"].update(
                lines=[""] * (MAX_ADDRESS_LINES + 1)
            ),
            "row_values": lambda context: context.update(
                lines=[{"values": [{}] * (MAX_COLUMNS + 1)}]
            ),
        }

        for name, mutate in cases.items():
            with self.subTest(name=name):
                context = self._editable_context()
                mutate(context)
                with self.assertRaisesRegex(EditableDocxError, "at most"):
                    render_editable_docx(context, "tr_TR")

    def test_editable_docx_never_silently_truncates_or_drops_input(self):
        context = self._editable_context()
        context["notes"] = ["X" * (MAX_TEXT_CHARS + 1)]
        with self.assertRaisesRegex(EditableDocxError, "character limit"):
            render_editable_docx(context, "tr_TR")

        context = self._editable_context()
        context["metadata"] = {"label": "Beklenmeyen yapı"}
        with self.assertRaisesRegex(EditableDocxError, "valid list"):
            render_editable_docx(context, "tr_TR")

    def test_editable_docx_applies_visibility_flags_independently(self):
        scenarios = (
            (
                {"show_company": False, "show_partner": True, "show_metadata": True},
                ("PARTNER-GÖRÜNÜR", "META-GÖRÜNÜR"),
                ("ŞİRKET-GÖRÜNÜR", "ŞİRKET-ADRESİ", "ŞİRKET-VKN-123"),
            ),
            (
                {"show_company": True, "show_partner": False, "show_metadata": False},
                ("ŞİRKET-GÖRÜNÜR", "ŞİRKET-ADRESİ", "ŞİRKET-VKN-123"),
                ("PARTNER-GÖRÜNÜR", "META-GÖRÜNÜR"),
            ),
            (
                {"show_company": True, "show_partner": True, "show_metadata": False},
                ("ŞİRKET-GÖRÜNÜR", "ŞİRKET-VKN-123", "PARTNER-GÖRÜNÜR"),
                ("META-GÖRÜNÜR",),
            ),
        )
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH

        for flags, visible, hidden in scenarios:
            with self.subTest(flags=flags):
                context = self._editable_context()
                context.update(flags)
                # A hidden company logo must not even be decoded.
                if not flags["show_company"]:
                    context["logo_bytes"] = b"not-an-image"
                content = render_editable_docx(context, "tr_TR")
                with ZipFile(BytesIO(content)) as archive:
                    searchable_xml = b"\n".join(
                        archive.read(name)
                        for name in archive.namelist()
                        if name == "word/document.xml" or name.startswith("word/header")
                    ).decode("utf-8")
                for marker in visible:
                    self.assertIn(marker, searchable_xml)
                for marker in hidden:
                    self.assertNotIn(marker, searchable_xml)

                if flags["show_metadata"]:
                    document = Document(BytesIO(content))
                    metadata_paragraph = next(
                        paragraph
                        for table in document.tables
                        for row in table.rows
                        for cell in row.cells
                        for paragraph in cell.paragraphs
                        if "META-GÖRÜNÜR" in paragraph.text
                    )
                    self.assertEqual(
                        metadata_paragraph.alignment,
                        WD_ALIGN_PARAGRAPH.CENTER,
                    )

    def test_editable_docx_accepts_decomposed_unicode_title_and_number(self):
        context = {
            "title": unicodedata.normalize("NFD", "Çağrı Belgesi"),
            "number": unicodedata.normalize("NFD", "İŞ-ÖÇĞ-1"),
            "company_info": {"name": "DocuCraft"},
            "partner_info": {"name": "Unicode Customer"},
            "columns": [],
            "lines": [],
            "totals": [],
        }
        content = render_editable_docx(context, "tr_TR")
        self.assertTrue(content.startswith(b"PK\x03\x04"))
        with ZipFile(BytesIO(content)) as archive:
            document_xml = archive.read("word/document.xml").decode("utf-8")
        self.assertIn("Çağrı Belgesi", document_xml)
        self.assertIn("İŞ-ÖÇĞ-1", document_xml)

    def test_editable_docx_export_uses_context_without_rendering_pdf(self):
        self.wizard.output_format = "docx_editable"
        record = self.env.user.partner_id
        context = {"title": "Teklif", "number": "S1"}
        with (
            patch.object(
                RdsExportWizard,
                "_record_context",
                return_value=context,
            ) as record_context,
            patch.object(RdsExportWizard, "_render_pdf") as render_pdf,
            patch(
                "odoo.addons.ranvals_document_studio.wizard.rds_export_wizard.render_editable_docx",
                return_value=b"PK\x03\x04editable",
            ) as render_word,
            patch(
                "odoo.addons.ranvals_document_studio.wizard.rds_localization.localize_docx",
                side_effect=lambda content, _language: content,
            ),
        ):
            outputs = self.wizard._export_record(record, "tr_TR")

        render_pdf.assert_not_called()
        record_context.assert_called_once_with(record, "tr_TR")
        render_word.assert_called_once_with(context, language_code="tr_TR")
        self.assertEqual(len(outputs), 1)
        self.assertTrue(outputs[0][0].endswith("_Editable_Word.docx"))
        self.assertEqual(outputs[0][2], DOCX_MIMETYPE)

    def test_editable_docx_applies_rtl_language_and_bounded_logo(self):
        logo_stream = BytesIO()
        Image.new("RGBA", (320, 120), (23, 52, 95, 255)).save(
            logo_stream,
            format="PNG",
        )
        context = {
            "title": "عرض سعر",
            "number": "S-AR-1",
            "direction": "rtl",
            "lang_code": "ar",
            "logo_bytes": logo_stream.getvalue(),
            "company_info": {"name": "شركة رانفالس"},
            "partner_info": {
                "name": "العميل",
                "email": "customer@example.com",
            },
            "columns": [
                {"label": "الوصف", "align": "left", "width": 65},
                {"label": "المبلغ", "align": "right", "width": 35},
            ],
            "lines": [
                {
                    "values": [
                        {"text": "خدمة استشارية", "align": "left"},
                        {"text": "1,234.50 USD", "align": "right"},
                    ]
                }
            ],
            "totals": [{"label": "الإجمالي", "value": "100", "is_total": True}],
            "bank": {"iban": "TR00 0000 0000", "swift": "ABCDEF12"},
            "bank_labels": {"iban": "IBAN", "swift": "SWIFT"},
        }

        content = render_editable_docx(context, "ar_001")

        with ZipFile(BytesIO(content)) as archive:
            document_xml = archive.read("word/document.xml")
            footer_xml = archive.read("word/footer1.xml")
            media = [name for name in archive.namelist() if name.startswith("word/media/")]
        self.assertIn(b"<w:bidi", document_xml)
        self.assertIn(b"<w:bidiVisual", document_xml)
        self.assertIn(b"ar-SA", document_xml)
        self.assertIn("صفحة".encode(), footer_xml)
        self.assertEqual(len(media), 1)
        self.assertNotIn(b"<wp:anchor", document_xml)

        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH

        document = Document(BytesIO(content))
        title_table = document.tables[0]
        # Logical cell order stays stable; bidiVisual performs the one and
        # only visual reversal in Word.
        self.assertIn("عرض سعر", title_table.cell(0, 0).text)
        self.assertIn("S-AR-1", title_table.cell(0, 1).text)
        number_run_xml = title_table.cell(0, 1).paragraphs[0].runs[0]._r.xml
        self.assertIn('w:rtl w:val="0"', number_run_xml)

        line_table = next(
            table
            for table in document.tables
            if table.rows and "الوصف" in table.cell(0, 0).text
        )
        self.assertEqual(
            line_table.cell(1, 0).paragraphs[0].alignment,
            WD_ALIGN_PARAGRAPH.RIGHT,
        )
        self.assertEqual(
            line_table.cell(1, 1).paragraphs[0].alignment,
            WD_ALIGN_PARAGRAPH.LEFT,
        )
        amount_run_xml = line_table.cell(1, 1).paragraphs[0].runs[0]._r.xml
        self.assertIn('w:rtl w:val="0"', amount_run_xml)
        iban_cell = next(
            cell
            for table in document.tables
            for row in table.rows
            for cell in row.cells
            if "TR00 0000 0000" in cell.text
        )
        iban_run = next(
            run
            for paragraph in iban_cell.paragraphs
            for run in paragraph.runs
            if "TR00 0000 0000" in run.text
        )
        self.assertIn('w:rtl w:val="0"', iban_run._r.xml)
        email_run = next(
            run
            for table in document.tables
            for row in table.rows
            for cell in row.cells
            for paragraph in cell.paragraphs
            for run in paragraph.runs
            if "customer@example.com" in run.text
        )
        self.assertIn('w:rtl w:val="0"', email_run._r.xml)

    def test_editable_docx_aspect_fits_wide_logo_and_reserves_header_space(self):
        logo_stream = BytesIO()
        Image.new("RGB", (8000, 40), (23, 52, 95)).save(
            logo_stream,
            format="PNG",
        )
        context = self._editable_context()
        context.update(
            {
                "logo_bytes": logo_stream.getvalue(),
                "logo_height_mm": 40,
            }
        )

        content = render_editable_docx(context, "tr_TR")

        with ZipFile(BytesIO(content)) as archive:
            safe_xml_parser = etree.XMLParser(resolve_entities=False, no_network=True)
            header_root = etree.fromstring(
                archive.read("word/header1.xml"), parser=safe_xml_parser
            )
            document_root = etree.fromstring(
                archive.read("word/document.xml"),
                parser=etree.XMLParser(resolve_entities=False, no_network=True),
            )
            settings_xml = archive.read("word/settings.xml")
            footer_xml = archive.read("word/footer1.xml")
        drawing_namespace = (
            "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
        )
        word_namespace = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
        extent = header_root.find(".//{%s}extent" % drawing_namespace)
        self.assertIsNotNone(extent)
        width_emu = int(extent.get("cx"))
        height_emu = int(extent.get("cy"))
        self.assertLessEqual(width_emu, 1_728_000)  # 4.8 cm
        self.assertLessEqual(height_emu, 1_440_000)  # 4.0 cm
        self.assertAlmostEqual(width_emu / height_emu, 200, delta=1)
        page_margins = document_root.find(".//{%s}pgMar" % word_namespace)
        self.assertGreaterEqual(int(page_margins.get("{%s}top" % word_namespace)), 2948)
        self.assertIn(b"updateFields", settings_xml)
        self.assertIn(b'val="true"', settings_xml)
        self.assertIn(b"PAGE", footer_xml)
        self.assertIn(b"NUMPAGES", footer_xml)
        self.assertGreaterEqual(footer_xml.count(b"<w:t>1</w:t>"), 2)

    def test_editable_docx_rejects_pathological_logo_before_pixel_decode(self):
        context = self._editable_context()
        context["logo_bytes"] = b"fake-image-header"
        image = MagicMock()
        image.size = (MAX_LOGO_DIMENSION + 1, 1)
        image.__enter__.return_value = image
        image.__exit__.return_value = False

        with (
            patch("PIL.Image.open", return_value=image),
            self.assertRaisesRegex(EditableDocxError, "dimensions exceed"),
        ):
            render_editable_docx(context, "tr_TR")

        image.load.assert_not_called()

    def test_editable_docx_converts_pillow_bomb_errors_to_safe_error(self):
        from PIL import Image as PillowImage

        context = self._editable_context()
        context["logo_bytes"] = b"fake-image-header"
        with (
            patch(
                "PIL.Image.open",
                side_effect=PillowImage.DecompressionBombError("bomb"),
            ),
            self.assertRaisesRegex(EditableDocxError, "valid, safe image"),
        ):
            render_editable_docx(context, "tr_TR")

    def test_editable_docx_supports_all_layout_profiles(self):
        from docx import Document
        from docx.oxml.ns import qn

        expected_fills = {
            "beauty": "17345F",
            "construction": "17345F",
            "technology": "17345F",
            "industrial": "D33A35",
            "eco": "17345F",
            "furniture": "D33A35",
            "noir_executive": "D33A35",
            "royal_ledger": "D33A35",
            "swiss_grid": "17345F",
            "arctic_minimal": "17345F",
            "indigo_flow": "D33A35",
            "emerald_ledger": "17345F",
            "sandstone_classic": "17345F",
            "graphite_copper": "D33A35",
        }
        fills = {}
        for layout_style in expected_fills:
            context = self._editable_context()
            context["layout_style"] = layout_style
            document = Document(
                BytesIO(render_editable_docx(context, "tr_TR"))
            )
            line_table = next(
                table
                for table in document.tables
                if table.rows and "Açıklama" in table.cell(0, 0).text
            )
            shading = line_table.cell(0, 0)._tc.get_or_add_tcPr().find(
                qn("w:shd")
            )
            fills[layout_style] = shading.get(qn("w:fill"))

        self.assertEqual(fills, expected_fills)

    def test_editable_docx_renders_all_business_connector_contexts(self):
        partner = self.env["res.partner"].create(
            {"name": "DocuCraft Native Word Partner", "lang": "en_US"}
        )
        sale = self.env["sale.order"].create({"partner_id": partner.id})
        purchase = self.env["purchase.order"].create({"partner_id": partner.id})
        journal = self.env["account.journal"].search(
            [("type", "=", "sale"), ("company_id", "=", self.env.company.id)],
            limit=1,
        )
        if not journal:
            journal = self.env["account.journal"].create(
                {
                    "name": "DocuCraft Native Word Journal",
                    "code": "DCNW",
                    "type": "sale",
                    "company_id": self.env.company.id,
                }
            )
        invoice = self.env["account.move"].create(
            {
                "move_type": "out_invoice",
                "partner_id": partner.id,
                "journal_id": journal.id,
                "invoice_date": fields.Date.context_today(self.env.user),
            }
        )

        from docx import Document

        for record in (sale, invoice, purchase):
            with self.subTest(model=record._name):
                template = self.env["rds.template"].get_default_for(
                    record._name,
                    company=record.company_id,
                )
                context = record._rds_document_context(template, "en_US")
                self.assertEqual(context["logo_height_mm"], template.logo_height_mm)
                self.assertEqual(context["layout_style"], template.layout_style)
                content = render_editable_docx(context, "en_US")
                document = Document(BytesIO(content))
                editable_text = "\n".join(
                    cell.text
                    for table in document.tables
                    for row in table.rows
                    for cell in row.cells
                )
                self.assertIn(context["title"], editable_text)
                self.assertIn(context["number"], editable_text)
                self.assertTrue(any(table.rows for table in document.tables))

    def test_editable_docx_rejects_unbounded_logo_and_invalid_context(self):
        with self.assertRaises(EditableDocxError):
            render_editable_docx([], "tr_TR")
        with self.assertRaises(EditableDocxError):
            render_editable_docx(
                {"title": "Teklif", "number": "S1", "logo_bytes": b"x" * (8 * 1024 * 1024 + 1)},
                "tr_TR",
            )

    def test_explicit_zip_contains_pdf_docx_and_png(self):
        self.wizard.output_format = "zip"
        expected_language = self.wizard.language_id.code or self.env.user.lang or "en_US"
        with (
            patch.object(
                RdsExportWizard, "_render_pdf", return_value=b"%PDF-test"
            ) as render_pdf,
            patch(
                "odoo.addons.ranvals_document_studio.wizard.rds_export_wizard.render_docx",
                return_value=b"PK\x03\x04docx",
            ) as render_word,
            patch.object(
                RdsExportWizard,
                "_record_context",
                return_value={"title": "Teklif", "number": "S1"},
            ),
            patch(
                "odoo.addons.ranvals_document_studio.wizard.rds_export_wizard.render_editable_docx",
                return_value=b"PK\x03\x04editable",
            ) as render_editable_word,
            patch(
                "odoo.addons.ranvals_document_studio.wizard.rds_export_wizard.render_png_pages",
                return_value=[b"PNG-page-1", b"PNG-page-2"],
            ),
            patch(
                "odoo.addons.ranvals_document_studio.wizard.rds_localization.localize_docx",
                side_effect=lambda content, _language: content,
            ) as localize_word,
        ):
            file_name, content, mimetype, record_map = self.wizard._build_output()

        render_pdf.assert_called_once()
        render_word.assert_called_once_with(b"%PDF-test", max_pages=50)
        render_editable_word.assert_called_once_with(
            {"title": "Teklif", "number": "S1"},
            language_code=expected_language,
        )
        self.assertEqual(localize_word.call_count, 2)
        localize_word.assert_any_call(b"PK\x03\x04docx", expected_language)
        localize_word.assert_any_call(b"PK\x03\x04editable", expected_language)
        self.assertTrue(file_name.endswith(".zip"))
        self.assertEqual(mimetype, "application/zip")
        self.assertEqual(len(record_map), 1)
        with ZipFile(BytesIO(content)) as archive:
            self.assertIsNone(archive.testzip())
            names = archive.namelist()
            self.assertEqual(len(names), 5)
            self.assertTrue(any(name.endswith(".pdf") for name in names))
            docx_names = [name for name in names if name.endswith(".docx")]
            self.assertEqual(len(docx_names), 2)
            self.assertTrue(any("Design_Preserved_Word" in name for name in docx_names))
            self.assertTrue(any("Editable_Word" in name for name in docx_names))
            self.assertEqual(sum(name.endswith(".png") for name in names), 2)

    def test_record_ids_are_strictly_validated(self):
        self.assertEqual(self.wizard._parse_record_ids("[2, 2, 1]"), [2, 1])
        for value in ('{"id": 1}', '["1"]', "[true]", "[0]", "[-1]", "[1.5]"):
            with self.subTest(value=value), self.assertRaises(UserError):
                self.wizard._parse_record_ids(value)

    def test_export_never_silently_drops_a_deleted_selection(self):
        partners = self.env["res.partner"].create([
            {"name": "DocuCraft complete selection A"},
            {"name": "DocuCraft complete selection B"},
        ])
        wizard = self.env["rds.export.wizard"].create({
            "res_model": "res.partner",
            "res_ids_json": "[%s, %s]" % tuple(partners.ids),
            "template_id": self.wizard.template_id.id,
            "output_format": "pdf",
        })
        partners[-1].unlink()

        with self.assertRaises(UserError):
            wizard._get_records()
        self.assertFalse(wizard._get_records(silent=True))

    def test_export_rejects_mixed_company_batch(self):
        foreign_company = self.env["res.company"].create({"name": "DocuCraft Mixed Export"})
        partners = self.env["res.partner"].sudo().create([
            {"name": "Current-company source", "company_id": self.env.company.id},
            {"name": "Foreign-company source", "company_id": foreign_company.id},
        ])
        wizard = self.env["rds.export.wizard"].sudo().create({
            "res_model": "res.partner",
            "res_ids_json": "[%s, %s]" % tuple(partners.ids),
            "template_id": self.wizard.template_id.id,
            "output_format": "pdf",
        })
        with self.assertRaises(ValidationError):
            wizard._build_output()

    def test_long_names_keep_record_id_suffixes_unique(self):
        records = self.env["res.partner"].create(
            [
                {"name": "A" * 180},
                {"name": "A" * 180},
            ]
        )
        self.wizard.output_format = "pdf"
        self.wizard.file_name_prefix = "P" * 180
        with patch.object(
            RdsExportWizard,
            "_render_pdf",
            return_value=b"%PDF-test",
        ):
            names = [
                self.wizard._export_record(record, "tr_TR")[0][0]
                for record in records
            ]

        self.assertEqual(len(names), len(set(names)))
        self.assertLessEqual(max(map(len, names)), 124)
        for record, name in zip(records, names, strict=True):
            self.assertIn("ID%s_" % record.id, name)

    def test_normal_user_can_store_result_but_not_another_users_wizard(self):
        owner = new_test_user(
            self.env,
            login="rds_wizard_owner",
            groups="base.group_user",
        )
        other_user = new_test_user(
            self.env,
            login="rds_wizard_other",
            groups="base.group_user",
        )
        wizard = self.env["rds.export.wizard"].with_user(owner).create(
            {
                "res_model": "res.partner",
                "res_ids_json": "[%s]" % owner.partner_id.id,
                "template_id": self.wizard.template_id.id,
                "output_format": "pdf",
            }
        )
        result = ("normal-user.pdf", b"%PDF-test", "application/pdf", [])
        with patch.object(RdsExportWizard, "_build_output", return_value=result):
            action = wizard.action_export()

        self.assertEqual(action["tag"], "ranvals_document_studio.download_export")
        self.assertEqual(
            base64.b64decode(wizard.sudo().file_data),
            b"%PDF-test",
        )
        with self.assertRaises(ValidationError):
            wizard.write({"res_ids_json": "[%s]" % self.env.user.partner_id.id})
        with (
            patch.object(RdsExportWizard, "_build_output", return_value=result),
            self.assertRaises(AccessError),
        ):
            wizard.with_user(other_user).action_export()

    def test_render_pdf_missing_connector_raises_user_error(self):
        with self.assertRaises(UserError):
            self.wizard._render_pdf(self.env.user.partner_id, "tr_TR")

    def test_output_budget_is_enforced_without_allocating_the_payload(self):
        with self.assertRaises(UserError):
            self.wizard._ensure_output_size(MAX_EXPORT_BYTES + 1)
