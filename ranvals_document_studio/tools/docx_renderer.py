import posixpath
from copy import deepcopy
from io import BytesIO
from zipfile import BadZipFile, ZipFile

from lxml import etree
from odoo import _

from .png_renderer import PngDependencyError, render_png_pages


DOCX_IMAGE_DPI = 300
DOCX_PAGE_WIDTH_CM = 21.0
DOCX_PAGE_HEIGHT_CM = 29.7
DOCX_RENDERER_FINGERPRINT = "DocuCraft PDF-faithful renderer v2 / 300 DPI"
MAX_DOCX_BYTES = 100 * 1024 * 1024


class DocxDependencyError(RuntimeError):
    pass


def _configure_image_quality(document, dpi=DOCX_IMAGE_DPI):
    """Disable Office image recompression and preserve the rendered page DPI."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    settings = document.settings.element
    if settings.find(qn("w:doNotAutoCompressPictures")) is None:
        settings.append(OxmlElement("w:doNotAutoCompressPictures"))
    default_dpi = settings.find(qn("w14:defaultImageDpi"))
    if default_dpi is None:
        default_dpi = OxmlElement("w14:defaultImageDpi")
        settings.append(default_dpi)
    default_dpi.set(qn("w14:val"), str(dpi))


def _configure_a4_page(section):
    from docx.shared import Cm

    section.page_width = Cm(DOCX_PAGE_WIDTH_CM)
    section.page_height = Cm(DOCX_PAGE_HEIGHT_CM)
    section.top_margin = Cm(0)
    section.bottom_margin = Cm(0)
    section.left_margin = Cm(0)
    section.right_margin = Cm(0)
    section.header_distance = Cm(0)
    section.footer_distance = Cm(0)
    section.gutter = Cm(0)


def _anchor_picture_to_page(inline_shape, page_number):
    """Turn an inline picture into a page-sized floating picture.

    A floating picture does not consume paragraph height, so the complete A4
    canvas can be preserved without producing a trailing blank Word page.
    """
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    inline = inline_shape._inline
    anchor = OxmlElement("wp:anchor")
    for name, value in (
        ("distT", "0"),
        ("distB", "0"),
        ("distL", "0"),
        ("distR", "0"),
        ("simplePos", "0"),
        ("relativeHeight", str(page_number)),
        ("behindDoc", "0"),
        ("locked", "1"),
        ("layoutInCell", "1"),
        ("allowOverlap", "1"),
    ):
        anchor.set(name, value)

    simple_position = OxmlElement("wp:simplePos")
    simple_position.set("x", "0")
    simple_position.set("y", "0")
    anchor.append(simple_position)

    for axis in ("H", "V"):
        position = OxmlElement("wp:position%s" % axis)
        position.set("relativeFrom", "page")
        offset = OxmlElement("wp:posOffset")
        offset.text = "0"
        position.append(offset)
        anchor.append(position)

    for tag in ("wp:extent", "wp:effectExtent"):
        element = inline.find(qn(tag))
        if element is not None:
            anchor.append(deepcopy(element))

    anchor.append(OxmlElement("wp:wrapNone"))
    for tag in ("wp:docPr", "wp:cNvGraphicFramePr", "a:graphic"):
        element = inline.find(qn(tag))
        if element is not None:
            anchor.append(deepcopy(element))

    doc_pr = anchor.find(qn("wp:docPr"))
    if doc_pr is not None:
        doc_pr.set("name", "PDF page %s" % page_number)
        doc_pr.set("descr", "PDF page %s" % page_number)

    inline.getparent().replace(inline, anchor)


def _add_pdf_page(document, page_png, page_number, is_last):
    from docx.enum.text import WD_BREAK
    from docx.shared import Cm, Pt

    paragraph = (
        document.paragraphs[0]
        if page_number == 1 and document.paragraphs
        else document.add_paragraph()
    )
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.left_indent = Cm(0)
    paragraph.paragraph_format.right_indent = Cm(0)
    paragraph.paragraph_format.line_spacing = Pt(1)

    picture = paragraph.add_run().add_picture(
        BytesIO(page_png),
        width=Cm(DOCX_PAGE_WIDTH_CM),
        height=Cm(DOCX_PAGE_HEIGHT_CM),
    )
    _anchor_picture_to_page(picture, page_number)
    if not is_last:
        paragraph.add_run().add_break(WD_BREAK.PAGE)


def _validate_pdf_faithful_archive(content, expected_pages):
    """Reject a malformed or accidentally regressed editable DOCX layout."""
    page_count = len(expected_pages)
    try:
        with ZipFile(BytesIO(content)) as archive:
            if archive.testzip() is not None:
                raise ValueError("corrupt OOXML member")
            document_xml = archive.read("word/document.xml")
            relationships_xml = archive.read("word/_rels/document.xml.rels")
            core_xml = archive.read("docProps/core.xml")
            media_payloads = {
                name: archive.read(name)
                for name in archive.namelist()
                if name.startswith("word/media/")
            }
        parser = etree.XMLParser(
            resolve_entities=False,
            no_network=True,
            load_dtd=False,
            huge_tree=False,
        )
        document_root = etree.fromstring(document_xml, parser=parser)
        relationships_root = etree.fromstring(relationships_xml, parser=parser)
    except (BadZipFile, etree.XMLSyntaxError, KeyError, ValueError) as exc:
        raise DocxDependencyError(
            _("The Word output was created, but the DOCX archive could not be validated.")
        ) from exc

    word_namespace = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    drawing_namespace = (
        "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
    )
    graphic_namespace = "http://schemas.openxmlformats.org/drawingml/2006/main"
    office_relationship_namespace = (
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    )
    package_relationship_namespace = (
        "http://schemas.openxmlformats.org/package/2006/relationships"
    )
    page_size = document_root.find(".//{%s}pgSz" % word_namespace)
    page_margins = document_root.find(".//{%s}pgMar" % word_namespace)
    anchors = document_root.findall(".//{%s}anchor" % drawing_namespace)
    page_geometry_is_valid = (
        page_size is not None
        and page_size.get("{%s}w" % word_namespace) == "11906"
        and page_size.get("{%s}h" % word_namespace) == "16838"
        and page_margins is not None
        and all(
            page_margins.get("{%s}%s" % (word_namespace, edge)) == "0"
            for edge in (
                "top",
                "right",
                "bottom",
                "left",
                "header",
                "footer",
                "gutter",
            )
        )
    )
    anchors_are_valid = len(anchors) == page_count and all(
        anchor.get("behindDoc") == "0"
        and anchor.find("{%s}wrapNone" % drawing_namespace) is not None
        and anchor.find("{%s}extent" % drawing_namespace) is not None
        and anchor.find("{%s}extent" % drawing_namespace).get("cx") == "7560000"
        and anchor.find("{%s}extent" % drawing_namespace).get("cy") == "10692000"
        and all(
            position is not None
            and position.get("relativeFrom") == "page"
            and position.find("{%s}posOffset" % drawing_namespace) is not None
            and position.find("{%s}posOffset" % drawing_namespace).text == "0"
            for position in (
                anchor.find("{%s}positionH" % drawing_namespace),
                anchor.find("{%s}positionV" % drawing_namespace),
            )
        )
        for anchor in anchors
    )

    relationships = {
        relation.get("Id"): relation.get("Target")
        for relation in relationships_root.findall(
            "{%s}Relationship" % package_relationship_namespace
        )
    }
    anchor_media = []
    referenced_media = set()
    for anchor in anchors:
        blip = anchor.find(".//{%s}blip" % graphic_namespace)
        relationship_id = (
            blip.get("{%s}embed" % office_relationship_namespace)
            if blip is not None
            else None
        )
        target = relationships.get(relationship_id)
        media_name = (
            posixpath.normpath(posixpath.join("word", target))
            if target
            else None
        )
        payload = media_payloads.get(media_name)
        if payload is None:
            anchors_are_valid = False
            break
        anchor_media.append(payload)
        referenced_media.add(media_name)

    from PIL import Image

    page_images_are_valid = True
    for payload in media_payloads.values():
        try:
            with Image.open(BytesIO(payload)) as image:
                image.load()
                dpi = image.info.get("dpi", (0, 0))
                valid_alpha = (
                    image.mode != "RGBA" or image.getchannel("A").getextrema() == (255, 255)
                )
                if not (
                    image.mode in ("RGB", "RGBA")
                    and image.width >= 2475
                    and image.height >= 3500
                    and abs(dpi[0] - DOCX_IMAGE_DPI) <= 1
                    and abs(dpi[1] - DOCX_IMAGE_DPI) <= 1
                    and valid_alpha
                ):
                    page_images_are_valid = False
                    break
        except (OSError, TypeError, ValueError):
            page_images_are_valid = False
            break

    if (
        not page_geometry_is_valid
        or not anchors_are_valid
        or not page_images_are_valid
        or anchor_media != list(expected_pages)
        or b"<w:tbl" in document_xml
        or referenced_media != set(media_payloads)
        or any(
            marker in document_xml
            for marker in (b"<a:duotone", b"<a:tint", b"<a:lumMod", b"<a:grayscl")
        )
        or DOCX_RENDERER_FINGERPRINT.encode("utf-8") not in core_xml
    ):
        raise DocxDependencyError(
            _(
                "The Word output could not be created with a page layout identical to the PDF. "
                "Update the DocuCraft core module and restart the Odoo workers."
            )
        )


def render_docx(pdf_bytes, max_pages=50):
    """Create a 300 DPI DOCX whose pages visually match the source PDF.

    The PDF is the single design source. Each PDF page is placed as a full-page
    A4 image in Word, intentionally prioritizing visual parity over editable
    document text.
    """
    if not isinstance(pdf_bytes, (bytes, bytearray)) or not pdf_bytes.startswith(
        b"%PDF"
    ):
        raise DocxDependencyError(
            _("A valid PDF document could not be created for Word output.")
        )
    if len(pdf_bytes) > MAX_DOCX_BYTES:
        raise DocxDependencyError(_("The Word output exceeds the 100 MB processing limit."))

    try:
        from docx import Document
    except ImportError as exc:  # pragma: no cover - depends on Odoo.sh requirements
        raise DocxDependencyError(
            _("The python-docx package is required for Word output. Add requirements.txt to the Odoo.sh root.")
        ) from exc

    try:
        pages = render_png_pages(
            bytes(pdf_bytes),
            dpi=DOCX_IMAGE_DPI,
            max_pages=max_pages,
            max_output_bytes=MAX_DOCX_BYTES,
        )
    except PngDependencyError as exc:
        raise DocxDependencyError(
            _("The PDF pages could not be prepared for Word output: %s") % exc
        ) from exc
    if not pages:
        raise DocxDependencyError(
            _("No usable pages were found in the PDF for Word output.")
        )

    document = Document()
    document.core_properties.comments = DOCX_RENDERER_FINGERPRINT
    document.core_properties.keywords = "PDF-faithful, full-page, 300-DPI"
    _configure_image_quality(document)
    _configure_a4_page(document.sections[0])
    document.styles["Normal"].paragraph_format.space_before = 0
    document.styles["Normal"].paragraph_format.space_after = 0

    for index, page_png in enumerate(pages, start=1):
        _add_pdf_page(document, page_png, index, index == len(pages))

    output = BytesIO()
    document.save(output)
    content = output.getvalue()
    if len(content) > MAX_DOCX_BYTES:
        raise DocxDependencyError(_("The Word output exceeds the 100 MB limit."))
    _validate_pdf_faithful_archive(content, pages)
    return content
