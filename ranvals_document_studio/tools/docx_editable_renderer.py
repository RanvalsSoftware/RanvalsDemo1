"""Native, editable DOCX renderer for DocuCraft document contexts.

Unlike :mod:`docx_renderer`, this module never rasterizes a PDF page.  It
builds an ordinary Word document whose text, tables, headings, notes and
totals remain editable.  The renderer deliberately accepts only the plain
mapping returned by ``_rds_document_context``; it performs no ORM lookups and
therefore cannot widen the caller's Odoo access rights.
"""

import re
import unicodedata
import warnings
from collections.abc import Mapping, Sequence
from io import BytesIO
from zipfile import BadZipFile, ZipFile

from lxml import etree
from odoo import _

EDITABLE_DOCX_FINGERPRINT = "DocuCraft native editable renderer v2"
MAX_EDITABLE_DOCX_BYTES = 100 * 1024 * 1024
MAX_LOGO_BYTES = 8 * 1024 * 1024
MAX_LOGO_PIXELS = 25_000_000
MAX_LOGO_DIMENSION = 12_000
MAX_TEXT_CHARS = 10_000
MAX_TOTAL_TEXT_CHARS = 1_000_000
MAX_COLUMNS = 24
MAX_LINES = 2_000
MAX_METADATA_ITEMS = 100
MAX_NOTES = 100
MAX_TOTALS = 100
MAX_ADDRESS_LINES = 10
MAX_LOGO_WIDTH_CM = 4.8
MAX_LOGO_HEIGHT_CM = 4.0

_A4_WIDTH_TWIPS = "11906"
_A4_HEIGHT_TWIPS = "16838"
_EMU_PER_CM = 360_000
_HEX_COLOR = re.compile(r"^#[0-9A-Fa-f]{6}$")
_CONTROL_CHARACTERS = re.compile(
    "[\x00-\x08\x0b\x0c\x0e-\x1f\x7f-\x9f]"
)
_LOCALES = {
    "tr": "tr-TR",
    "tr_TR": "tr-TR",
    "en": "en-US",
    "en_GB": "en-GB",
    "en_US": "en-US",
    "ru": "ru-RU",
    "ru_RU": "ru-RU",
    "fr": "fr-FR",
    "fr_FR": "fr-FR",
    "de": "de-DE",
    "de_DE": "de-DE",
    "es": "es-ES",
    "es_ES": "es-ES",
    "it": "it-IT",
    "it_IT": "it-IT",
    "pt": "pt-PT",
    "pt_BR": "pt-BR",
    "pt_PT": "pt-PT",
    "ar": "ar-SA",
    "ar_001": "ar-SA",
}
_PAGE_LABELS = {
    "tr": "Sayfa",
    "en": "Page",
    "ru": "Страница",
    "fr": "Page",
    "de": "Seite",
    "es": "Página",
    "it": "Pagina",
    "pt": "Página",
    "ar": "صفحة",
}


def _resolve_docx_locale(language_code):
    """Return a Word proofing locale for every supported Odoo locale variant."""
    code = str(language_code or "en_US")
    if code in _LOCALES:
        return _LOCALES[code]
    prefix = code.split("_", 1)[0].lower()
    return _LOCALES.get(prefix, "en-US")


_DOCX_FONT_KEYS = {
    "serif": "Georgia",
    "sans": "Arial",
    "technical": "Trebuchet MS",
    "editorial": "Georgia",
    "lato": "Lato",
    "roboto": "Roboto",
    "open_sans": "Open Sans",
    "montserrat": "Montserrat",
    "raleway": "Raleway",
    "oswald": "Oswald",
    "tajawal": "Tajawal",
    "fira_mono": "Fira Mono",
    "humanist": "Calibri",
    "helvetica": "Arial",
    "verdana": "Verdana",
    "tahoma": "Tahoma",
    "lucida": "Lucida Sans",
    "times": "Times New Roman",
    "garamond": "Garamond",
    "palatino": "Palatino Linotype",
    "cambria": "Cambria",
    "bookman": "Bookman Old Style",
    "monospace": "Courier New",
}
_DOCX_CSS_FAMILIES = {
    "georgia": "Georgia",
    "arial": "Arial",
    "trebuchet ms": "Trebuchet MS",
    "lato": "Lato",
    "roboto": "Roboto",
    "open_sans": "Open Sans",
    "open sans": "Open Sans",
    "montserrat": "Montserrat",
    "raleway": "Raleway",
    "oswald": "Oswald",
    "tajawal": "Tajawal",
    "fira_mono": "Fira Mono",
    "fira mono": "Fira Mono",
    "calibri": "Calibri",
    "helvetica": "Arial",
    "verdana": "Verdana",
    "tahoma": "Tahoma",
    "lucida sans": "Lucida Sans",
    "times new roman": "Times New Roman",
    "garamond": "Garamond",
    "palatino linotype": "Palatino Linotype",
    "cambria": "Cambria",
    "bookman old style": "Bookman Old Style",
    "courier new": "Courier New",
}

# Native Word keeps one robust, editable document structure while each QWeb
# layout supplies a deliberate colour hierarchy.  Values refer to keys in the
# resolved template theme; keeping the profile declarative makes new styles
# explicit and prevents an unknown layout from silently adopting an arbitrary
# appearance.
_LAYOUT_PROFILES = {
    "beauty": ("accent", "primary", "primary"),
    "construction": ("primary", "primary", "primary"),
    "technology": ("primary", "primary", "primary"),
    "industrial": ("primary", "accent", "accent"),
    "eco": ("primary", "primary", "primary"),
    "furniture": ("accent", "accent", "primary"),
    "noir_executive": ("primary", "accent", "accent"),
    "royal_ledger": ("primary", "accent", "primary"),
    "swiss_grid": ("accent", "primary", "accent"),
    "arctic_minimal": ("accent", "primary", "primary"),
    "indigo_flow": ("primary", "accent", "primary"),
    "emerald_ledger": ("primary", "primary", "accent"),
    "sandstone_classic": ("accent", "primary", "accent"),
    "graphite_copper": ("primary", "accent", "accent"),
}


class EditableDocxError(RuntimeError):
    """Raised when a native DOCX cannot be produced or validated safely."""


class _TextBudget:
    def __init__(self):
        self.used = 0

    def clean(self, value, *, limit=MAX_TEXT_CHARS):
        if value is None or value is False:
            return ""
        if isinstance(value, bytes):
            try:
                value = value.decode("utf-8")
            except UnicodeDecodeError as exc:
                raise EditableDocxError(
                    _("The editable Word document contains invalid UTF-8 text.")
                ) from exc
        elif not isinstance(value, str):
            value = str(value)
        value = unicodedata.normalize("NFC", _CONTROL_CHARACTERS.sub("", value))
        if len(value) > limit:
            raise EditableDocxError(
                _("A text value in the editable Word document exceeds the %s character limit.")
                % limit
            )
        self.used += len(value)
        if self.used > MAX_TOTAL_TEXT_CHARS:
            raise EditableDocxError(
                _("The total text in the editable Word document exceeds the safe processing limit.")
            )
        return value


def _mapping(value):
    return value if isinstance(value, Mapping) else {}


def _items(value, maximum, label):
    if value is None or value is False:
        return []
    if isinstance(value, Sequence) and not isinstance(value, (str, bytes, bytearray)):
        if len(value) > maximum:
            raise EditableDocxError(
                _("%s can contain at most %s items.") % (label, maximum)
            )
        return list(value)
    raise EditableDocxError(_("%s must be a valid list.") % label)


def _color(value, fallback):
    return value[1:].upper() if isinstance(value, str) and _HEX_COLOR.fullmatch(value) else fallback


def _font(value, fallback="Arial"):
    """Resolve an allow-listed template key or CSS stack to a Word font name."""
    if not isinstance(value, str):
        return fallback
    normalized = value.strip().casefold()
    if normalized in _DOCX_FONT_KEYS:
        return _DOCX_FONT_KEYS[normalized]
    first_family = value.split(",", 1)[0].strip().strip("'\"").casefold()
    return _DOCX_CSS_FAMILIES.get(first_family, fallback)


def _positive_number(value, fallback=0.0):
    try:
        result = float(value)
    except (TypeError, ValueError, OverflowError):
        return fallback
    return result if 0.0 <= result <= 10_000.0 else fallback


def _text_prefers_rtl(value):
    """Return the direction of the first strong Unicode character.

    Numbers, e-mail addresses, URLs, document numbers, SWIFT codes and IBANs
    deliberately remain LTR inside an RTL paragraph.  This prevents Word from
    visually reordering punctuation and account/document identifiers.
    """
    for character in value or "":
        direction = unicodedata.bidirectional(character)
        if direction in ("R", "AL"):
            return True
        if direction == "L":
            return False
    return False


def _set_cell_shading(cell, fill):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)
    shading.set(qn("w:val"), "clear")


def _set_cell_margins(cell, *, top=80, start=100, bottom=80, end=100):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn("w:%s" % edge))
        if node is None:
            node = OxmlElement("w:%s" % edge)
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _set_table_direction(table, rtl):
    if not rtl:
        return
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    properties = table._tbl.tblPr
    bidi = properties.find(qn("w:bidiVisual"))
    if bidi is None:
        bidi = OxmlElement("w:bidiVisual")
        properties.append(bidi)
    bidi.set(qn("w:val"), "1")


def _set_paragraph_direction(paragraph, rtl, alignment=None):
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    alignments = {
        "left": WD_ALIGN_PARAGRAPH.RIGHT if rtl else WD_ALIGN_PARAGRAPH.LEFT,
        "right": WD_ALIGN_PARAGRAPH.LEFT if rtl else WD_ALIGN_PARAGRAPH.RIGHT,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
    }
    if alignment in alignments:
        paragraph.alignment = alignments[alignment]
    elif rtl:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    if rtl:
        properties = paragraph._p.get_or_add_pPr()
        bidi = properties.find(qn("w:bidi"))
        if bidi is None:
            bidi = OxmlElement("w:bidi")
            properties.append(bidi)
        bidi.set(qn("w:val"), "1")


def _set_ooxml_font_slots(properties, font):
    """Set the same family for every Word script slot."""
    from docx.oxml.ns import qn

    fonts = properties.get_or_add_rFonts()
    for slot in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn("w:%s" % slot), font)


def _set_run_font(run, font):
    run.font.name = font
    _set_ooxml_font_slots(run._r.get_or_add_rPr(), font)


def _set_style_font(style, font):
    style.font.name = font
    _set_ooxml_font_slots(style._element.get_or_add_rPr(), font)


def _set_run_language(run, locale, rtl=False, explicit_direction=False):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    properties = run._r.get_or_add_rPr()
    language = properties.find(qn("w:lang"))
    if language is None:
        language = OxmlElement("w:lang")
        properties.append(language)
    language.set(qn("w:val"), locale)
    if explicit_direction:
        rtl_node = properties.find(qn("w:rtl"))
        if rtl_node is None:
            rtl_node = OxmlElement("w:rtl")
            properties.append(rtl_node)
        rtl_node.set(qn("w:val"), "1" if rtl else "0")
    if rtl:
        language.set(qn("w:bidi"), locale)


def _add_text(
    paragraph,
    text,
    *,
    font,
    size,
    color,
    locale,
    rtl=False,
    bold=False,
    italic=False,
    alignment=None,
):
    from docx.shared import Pt, RGBColor

    _set_paragraph_direction(paragraph, rtl, alignment)
    run = paragraph.add_run(text)
    run.bold = bool(bold)
    run.italic = bool(italic)
    _set_run_font(run, font)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run_rtl = rtl and _text_prefers_rtl(text)
    _set_run_language(
        run,
        locale,
        rtl=run_rtl,
        explicit_direction=rtl,
    )
    return run


def _prepare_logo(payload):
    if not payload:
        return b"", 0, 0
    if not isinstance(payload, (bytes, bytearray)) or len(payload) > MAX_LOGO_BYTES:
        raise EditableDocxError(_("The company logo exceeds the 8 MB safe processing limit."))
    try:
        from PIL import Image, ImageOps
    except ImportError as exc:
        raise EditableDocxError(
            _("The Pillow package is required to process the company logo.")
        ) from exc

    try:
        source = BytesIO(bytes(payload))
        with warnings.catch_warnings():
            warnings.simplefilter("error", Image.DecompressionBombWarning)
            image = Image.open(source)
        with image:
            # Validate dimensions from the image header before decoding pixel
            # data.  This stops decompression bombs and pathological 1xN/Nx1
            # images before ``load`` allocates their canvas.
            width, height = image.size
            if (
                width <= 0
                or height <= 0
                or width > MAX_LOGO_DIMENSION
                or height > MAX_LOGO_DIMENSION
                or width * height > MAX_LOGO_PIXELS
            ):
                raise EditableDocxError(
                    _("The company logo dimensions exceed the safe processing limit.")
                )
            image.load()
            image = ImageOps.exif_transpose(image)
            image.thumbnail((1600, 600))
            if image.mode not in ("RGB", "RGBA"):
                image = image.convert("RGBA" if "A" in image.getbands() else "RGB")
            sanitized_width, sanitized_height = image.size
            output = BytesIO()
            image.save(output, format="PNG", optimize=True)
            sanitized = output.getvalue()
    except EditableDocxError:
        raise
    except (
        Image.DecompressionBombError,
        Image.DecompressionBombWarning,
        OSError,
        ValueError,
        SyntaxError,
    ) as exc:
        raise EditableDocxError(_("The company logo is not a valid, safe image.")) from exc
    if len(sanitized) > MAX_LOGO_BYTES:
        raise EditableDocxError(_("The processed company logo exceeds the 8 MB limit."))
    return sanitized, sanitized_width, sanitized_height


def _logo_extent(width_px, height_px, requested_height_mm):
    requested_height_cm = min(
        max(_positive_number(requested_height_mm, fallback=18.0) / 10.0, 0.8),
        MAX_LOGO_HEIGHT_CM,
    )
    width_cm = requested_height_cm * width_px / height_px
    scale = min(
        1.0,
        MAX_LOGO_WIDTH_CM / width_cm if width_cm else 1.0,
        MAX_LOGO_HEIGHT_CM / requested_height_cm,
    )
    return max(width_cm * scale, 0.01), max(requested_height_cm * scale, 0.01)


def _set_repeat_table_header(row):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    properties = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    properties.append(repeat)


def _prevent_row_split(row):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    properties = row._tr.get_or_add_trPr()
    if properties.find(qn("w:cantSplit")) is None:
        properties.append(OxmlElement("w:cantSplit"))


def _remove_cell_paragraph(cell):
    paragraph = cell.paragraphs[0]
    for child in list(paragraph._p):
        paragraph._p.remove(child)
    return paragraph


def _add_info_lines(
    cell,
    info,
    budget,
    *,
    font,
    color,
    locale,
    rtl,
    labels=None,
    tax_label=None,
):
    info = _mapping(info)
    labels = _mapping(labels)
    values = []
    name = budget.clean(info.get("name"))
    if name:
        values.append((name, True))
    for item in _items(info.get("lines"), MAX_ADDRESS_LINES, _("Address lines")):
        text = budget.clean(item)
        if text:
            values.append((text, False))
    contact_labels = (
        ("phone", budget.clean(labels.get("phone")) or "Phone"),
        ("email", budget.clean(labels.get("email")) or "Email"),
        ("website", budget.clean(labels.get("website")) or "Website"),
        (
            "vat",
            tax_label
            or budget.clean(labels.get("tax_no"))
            or "Tax ID",
        ),
    )
    for key, prefix in contact_labels:
        text = budget.clean(info.get(key))
        if text:
            values.append(("%s: %s" % (prefix, text), False))
    if not values:
        values.append(("-", False))
    for text, bold in values:
        # The first cell paragraph belongs to the block heading.  Reusing it
        # made e.g. ``COMPANY INFORMATIONMy Company`` appear as one run in
        # Word.  Every value therefore gets its own editable paragraph.
        paragraph = cell.add_paragraph()
        paragraph.paragraph_format.space_after = 0
        _add_text(
            paragraph,
            text,
            font=font,
            size=9 if not bold else 10,
            color=color,
            locale=locale,
            rtl=rtl,
            bold=bold,
            alignment="left",
        )


def _configure_document(
    document,
    *,
    body_font,
    heading_font,
    locale,
    rtl,
    has_company_header,
):
    from docx.enum.style import WD_STYLE_TYPE
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt

    section = document.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    # Reserve the complete maximum logo/header canvas above the body.  Word
    # otherwise lets a tall logo or wrapped company details overlap the first
    # body table because headers live inside the top margin.
    section.top_margin = Cm(5.2 if has_company_header else 1.65)
    section.bottom_margin = Cm(1.65)
    section.left_margin = Cm(1.55)
    section.right_margin = Cm(1.55)
    section.header_distance = Cm(0.65)
    section.footer_distance = Cm(0.65)

    normal = document.styles["Normal"]
    _set_style_font(normal, body_font)
    normal.font.size = Pt(9)
    normal.paragraph_format.space_after = Pt(3)

    for style_name, size in (("Title", 24), ("Heading 1", 15), ("Heading 2", 11)):
        style = document.styles[style_name]
        _set_style_font(style, heading_font)
        style.font.size = Pt(size)

    if "DocuCraft Small" not in document.styles:
        small = document.styles.add_style("DocuCraft Small", WD_STYLE_TYPE.PARAGRAPH)
    else:
        small = document.styles["DocuCraft Small"]
    _set_style_font(small, body_font)
    small.font.size = Pt(8)
    document.core_properties.comments = EDITABLE_DOCX_FINGERPRINT
    document.core_properties.keywords = "native-editable, tables, text"
    document.core_properties.subject = "Editable business document"
    settings = document.settings.element
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "true")
    return section


def _add_page_field(paragraph, instruction, *, font, locale, rtl):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    run = paragraph.add_run()
    _set_run_font(run, font)
    _set_run_language(run, locale, rtl=rtl)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    text = OxmlElement("w:instrText")
    text.set(qn("xml:space"), "preserve")
    text.text = " %s " % instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    cached = OxmlElement("w:t")
    # Non-Word previewers often do not evaluate fields.  A safe cached result
    # keeps the footer useful there, while ``w:updateFields`` refreshes the
    # actual PAGE/NUMPAGES values when Word opens the file.
    cached.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for node in (begin, text, separate, cached, end):
        run._r.append(node)


def _render_header(section, context, budget, style, logo):
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm

    if not bool(context.get("show_company", True)):
        return

    header = section.header
    table = header.add_table(rows=1, cols=2, width=Cm(17.9))
    table.autofit = False
    table.columns[0].width = Cm(5.1)
    table.columns[1].width = Cm(12.8)
    _set_table_direction(table, style["rtl"])
    # Keep logical cell order stable. ``w:bidiVisual`` alone mirrors the
    # presentation in RTL documents; manually swapping these cells as well
    # would reverse them twice.
    logo_cell, info_cell = table.rows[0].cells
    for cell in (logo_cell, info_cell):
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        _set_cell_margins(cell, top=0, bottom=20, start=0, end=0)

    logo_bytes, logo_width_px, logo_height_px = logo
    if logo_bytes:
        paragraph = _remove_cell_paragraph(logo_cell)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT if style["rtl"] else WD_ALIGN_PARAGRAPH.LEFT
        width_cm, height_cm = _logo_extent(
            logo_width_px,
            logo_height_px,
            context.get("logo_height_mm"),
        )
        paragraph.add_run().add_picture(
            BytesIO(logo_bytes),
            width=Cm(width_cm),
            height=Cm(height_cm),
        )

    company_info = _mapping(context.get("company_info"))
    company_name = budget.clean(company_info.get("name"))
    paragraph = _remove_cell_paragraph(info_cell)
    paragraph.paragraph_format.space_after = 0
    _add_text(
        paragraph,
        company_name or " ",
        font=style["heading_font"],
        size=11,
        color=style["primary"],
        locale=style["locale"],
        rtl=style["rtl"],
        bold=True,
        alignment="right",
    )
    tagline = budget.clean(context.get("tagline"))
    if tagline:
        paragraph = info_cell.add_paragraph()
        paragraph.paragraph_format.space_after = 0
        _add_text(
            paragraph,
            tagline,
            font=style["body_font"],
            size=8,
            color=style["text"],
            locale=style["locale"],
            rtl=style["rtl"],
            alignment="right",
        )
    company_details = []
    for item in _items(
        company_info.get("lines"),
        MAX_ADDRESS_LINES,
        _("Company address lines"),
    ):
        detail = budget.clean(item)
        if detail:
            company_details.append(detail)
    labels = _mapping(context.get("labels"))
    contact_labels = {
        "phone": budget.clean(labels.get("phone")) or "Phone",
        "email": budget.clean(labels.get("email")) or "Email",
        "website": budget.clean(labels.get("website")) or "Website",
        "vat": (
            budget.clean(context.get("tax_label"))
            or budget.clean(labels.get("tax_no"))
            or "Tax ID"
        ),
    }
    for key, label in contact_labels.items():
        detail = budget.clean(company_info.get(key))
        if detail:
            company_details.append("%s: %s" % (label, detail))
    if company_details:
        paragraph = info_cell.add_paragraph()
        paragraph.paragraph_format.space_after = 0
        _add_text(
            paragraph,
            " • ".join(company_details),
            font=style["body_font"],
            size=7.5,
            color=style["text"],
            locale=style["locale"],
            rtl=style["rtl"],
            alignment="right",
        )


def _render_footer(section, context, budget, style):
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    footer = section.footer
    paragraph = footer.paragraphs[0]
    paragraph.paragraph_format.space_before = 0
    paragraph.paragraph_format.space_after = 0
    footer_text = budget.clean(context.get("footer_text")) if context.get("show_footer", True) else ""
    if footer_text:
        _add_text(
            paragraph,
            footer_text,
            font=style["body_font"],
            size=8,
            color=style["text"],
            locale=style["locale"],
            rtl=style["rtl"],
            alignment="center",
        )
        separator = paragraph.add_run("  •  ")
        _set_run_font(separator, style["body_font"])
        _set_run_language(separator, style["locale"], rtl=style["rtl"])
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    page_label = paragraph.add_run(
        "%s " % _PAGE_LABELS.get(style["language_prefix"], "Page")
    )
    _set_run_font(page_label, style["body_font"])
    _set_run_language(page_label, style["locale"], rtl=style["rtl"])
    _add_page_field(
        paragraph,
        "PAGE",
        font=style["body_font"],
        locale=style["locale"],
        rtl=style["rtl"],
    )
    separator = paragraph.add_run(" / ")
    _set_run_font(separator, style["body_font"])
    _set_run_language(separator, style["locale"], rtl=style["rtl"])
    _add_page_field(
        paragraph,
        "NUMPAGES",
        font=style["body_font"],
        locale=style["locale"],
        rtl=style["rtl"],
    )
    _set_paragraph_direction(paragraph, style["rtl"], "center")


def _render_title(document, context, budget, style):
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
    from docx.shared import Cm, Pt

    table = document.add_table(rows=1, cols=2)
    table.autofit = False
    table.columns[0].width = Cm(12.5)
    table.columns[1].width = Cm(5.4)
    _set_table_direction(table, style["rtl"])
    # Logical order is title then number for both directions.  Word mirrors
    # the visual order through ``w:bidiVisual`` in RTL mode.
    title_cell, number_cell = table.rows[0].cells
    for cell in (title_cell, number_cell):
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        _set_cell_margins(cell, top=160, bottom=160, start=140, end=140)
    _set_cell_shading(title_cell, style["title_fill"])
    _set_cell_shading(number_cell, style["secondary"])

    title = budget.clean(context.get("title")) or _("Document")
    paragraph = _remove_cell_paragraph(title_cell)
    paragraph.paragraph_format.space_after = 0
    _add_text(
        paragraph,
        title,
        font=style["heading_font"],
        size=18,
        color="FFFFFF",
        locale=style["locale"],
        rtl=style["rtl"],
        bold=True,
        alignment="left",
    )
    number = budget.clean(context.get("number")) or "-"
    paragraph = _remove_cell_paragraph(number_cell)
    paragraph.paragraph_format.space_after = 0
    _add_text(
        paragraph,
        number,
        font=style["heading_font"],
        size=13,
        color=style["accent"],
        locale=style["locale"],
        rtl=style["rtl"],
        bold=True,
        alignment="right",
    )
    spacer = document.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)


def _render_party_and_metadata(document, context, budget, style):
    from docx.shared import Cm, Pt

    show_partner = bool(context.get("show_partner", True))
    show_company = bool(context.get("show_company", True))
    show_metadata = bool(context.get("show_metadata", True))
    metadata = (
        _items(context.get("metadata"), MAX_METADATA_ITEMS, _("Document information"))
        if show_metadata
        else []
    )
    labels = _mapping(context.get("labels"))
    blocks = []
    if show_company:
        blocks.append(
            (
                "info",
                budget.clean(labels.get("company_info")) or _("Company Information"),
                context.get("company_info"),
            )
        )
    if show_partner:
        blocks.append(
            (
                "info",
                budget.clean(labels.get("partner_info"))
                or _("Customer / Vendor"),
                context.get("partner_info"),
            )
        )
    if show_metadata:
        blocks.append(
            (
                "metadata",
                budget.clean(labels.get("document_info")) or _("Document Information"),
                metadata,
            )
        )
    if not blocks:
        return

    table = document.add_table(rows=1, cols=len(blocks))
    table.autofit = False
    block_width = 17.9 / len(blocks)
    for column in table.columns:
        column.width = Cm(block_width)
    _set_table_direction(table, style["rtl"])
    for cell in table.rows[0].cells:
        _set_cell_margins(cell, top=100, bottom=100, start=120, end=120)

    for cell, (kind, heading, payload) in zip(
        table.rows[0].cells,
        blocks,
        strict=True,
    ):
        paragraph = _remove_cell_paragraph(cell)
        paragraph.paragraph_format.space_after = Pt(2)
        _add_text(
            paragraph,
            heading,
            font=style["heading_font"],
            size=10,
            color=style["primary"],
            locale=style["locale"],
            rtl=style["rtl"],
            bold=True,
            alignment="left",
        )
        if kind == "info":
            _add_info_lines(
                cell,
                payload,
                budget,
                font=style["body_font"],
                color=style["text"],
                locale=style["locale"],
                rtl=style["rtl"],
                labels=labels,
                tax_label=budget.clean(context.get("tax_label")),
            )
            continue

        for item in payload:
            item = _mapping(item)
            label = budget.clean(item.get("label"))
            value = budget.clean(item.get("value"))
            if not label and not value:
                continue
            alignment = item.get("align")
            if alignment not in ("left", "right", "center"):
                alignment = "left"
            paragraph = cell.add_paragraph()
            paragraph.paragraph_format.space_after = 0
            _add_text(
                paragraph,
                "%s: " % label if label else "",
                font=style["body_font"],
                size=8.5,
                color=style["text"],
                locale=style["locale"],
                rtl=style["rtl"],
                bold=True,
                alignment=alignment,
            )
            _add_text(
                paragraph,
                value or "-",
                font=style["body_font"],
                size=8.5,
                color=style["text"],
                locale=style["locale"],
                rtl=style["rtl"],
            )
    document.add_paragraph().paragraph_format.space_after = 0


def _render_lines(document, context, budget, style):
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
    from docx.shared import Cm

    columns = [
        _mapping(item)
        for item in _items(context.get("columns"), MAX_COLUMNS, _("Document columns"))
    ]
    if not columns:
        columns = [{"label": _("Description"), "align": "left", "width": 100}]
    lines = [
        _mapping(item)
        for item in _items(context.get("lines"), MAX_LINES, _("Document lines"))
    ]
    table = document.add_table(rows=1, cols=len(columns))
    table.style = "Table Grid"
    table.autofit = False
    _set_table_direction(table, style["rtl"])
    header = table.rows[0]
    _set_repeat_table_header(header)
    widths = []
    total_width = sum(_positive_number(item.get("width")) for item in columns)
    for item in columns:
        raw_width = _positive_number(item.get("width"))
        widths.append(17.9 * (raw_width / total_width if total_width else 1 / len(columns)))
    for index, (cell, column, width) in enumerate(zip(header.cells, columns, widths, strict=True)):
        cell.width = Cm(width)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        _set_cell_shading(cell, style["line_header_fill"])
        _set_cell_margins(cell, start=60, end=60)
        paragraph = _remove_cell_paragraph(cell)
        label = budget.clean(column.get("label")) or " "
        # Compact single-word headers such as DISCOUNT must remain readable in
        # intentionally narrow numeric columns instead of breaking one final
        # letter onto a separate line in Word.
        header_size = 7 if width < 2.0 and len(label) > 7 else 8
        _add_text(
            paragraph,
            label,
            font=style["heading_font"],
            size=header_size,
            color="FFFFFF",
            locale=style["locale"],
            rtl=style["rtl"],
            bold=True,
            alignment=column.get("align") if column.get("align") in ("left", "right", "center") else "left",
        )

    for line_index, line in enumerate(lines):
        row = table.add_row()
        _prevent_row_split(row)
        if line.get("is_section") or line.get("is_note"):
            merged = row.cells[0].merge(row.cells[-1])
            _set_cell_margins(merged)
            if line.get("is_section"):
                _set_cell_shading(merged, style["secondary"])
            paragraph = _remove_cell_paragraph(merged)
            _add_text(
                paragraph,
                budget.clean(line.get("description")) or " ",
                font=style["heading_font"] if line.get("is_section") else style["body_font"],
                size=9,
                color=style["primary"] if line.get("is_section") else style["text"],
                locale=style["locale"],
                rtl=style["rtl"],
                bold=bool(line.get("is_section")),
                italic=bool(line.get("is_note")),
                alignment="left",
            )
            continue

        values = [
            _mapping(value)
            for value in _items(
                line.get("values"),
                MAX_COLUMNS,
                _("Line cells"),
            )
        ]
        for index, cell in enumerate(row.cells):
            cell.width = Cm(widths[index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            _set_cell_margins(cell)
            if line_index % 2:
                _set_cell_shading(cell, "FAFBFC")
            value = values[index] if index < len(values) else {}
            if value.get("highlight"):
                _set_cell_shading(cell, style["secondary"])
            paragraph = _remove_cell_paragraph(cell)
            alignment = value.get("align") or columns[index].get("align") or "left"
            _add_text(
                paragraph,
                budget.clean(value.get("text")) or " ",
                font=style["body_font"],
                size=8.5,
                color=style["text"],
                locale=style["locale"],
                rtl=style["rtl"],
                bold=bool(value.get("bold")),
                alignment=alignment if alignment in ("left", "right", "center") else "left",
            )
    document.add_paragraph().paragraph_format.space_after = 0


def _render_totals(document, context, budget, style):
    from docx.shared import Cm

    totals = [
        _mapping(item)
        for item in _items(context.get("totals"), MAX_TOTALS, _("Document totals"))
    ]
    if not totals:
        return
    table = document.add_table(rows=0, cols=2)
    table.autofit = False
    table.alignment = 2 if not style["rtl"] else 0
    table.columns[0].width = Cm(4.7)
    table.columns[1].width = Cm(4.1)
    _set_table_direction(table, style["rtl"])
    for item in totals:
        row = table.add_row()
        _prevent_row_split(row)
        is_total = bool(item.get("is_total"))
        for cell in row.cells:
            _set_cell_margins(cell)
            if is_total:
                _set_cell_shading(cell, style["total_fill"])
        label = budget.clean(item.get("label")) or " "
        value = budget.clean(item.get("value")) or " "
        for cell, text, alignment in (
            (row.cells[0], label, "left"),
            (row.cells[1], value, "right"),
        ):
            paragraph = _remove_cell_paragraph(cell)
            _add_text(
                paragraph,
                text,
                font=style["heading_font"] if is_total else style["body_font"],
                size=10 if is_total else 9,
                color="FFFFFF" if is_total else style["text"],
                locale=style["locale"],
                rtl=style["rtl"],
                bold=is_total,
                alignment=alignment,
            )


def _render_notes_and_bank(document, context, budget, style):
    from docx.shared import Pt

    labels = _mapping(context.get("labels"))
    notes = (
        _items(context.get("notes"), MAX_NOTES, _("Document notes"))
        if context.get("show_notes", True)
        else []
    )
    if notes:
        heading = document.add_paragraph()
        heading.paragraph_format.space_before = Pt(8)
        _add_text(
            heading,
            budget.clean(labels.get("notes")) or _("Notes / Payment Terms"),
            font=style["heading_font"],
            size=10,
            color=style["primary"],
            locale=style["locale"],
            rtl=style["rtl"],
            bold=True,
            alignment="left",
        )
        for item in notes:
            text = budget.clean(item)
            if not text:
                continue
            paragraph = document.add_paragraph(style="List Bullet")
            paragraph.paragraph_format.space_after = 0
            _add_text(
                paragraph,
                text,
                font=style["body_font"],
                size=8.5,
                color=style["text"],
                locale=style["locale"],
                rtl=style["rtl"],
                alignment="left",
            )

    bank = _mapping(context.get("bank")) if context.get("show_bank", True) else {}
    bank_labels = _mapping(context.get("bank_labels"))
    bank_values = []
    for key in ("bank_name", "branch", "account_name", "iban", "swift"):
        value = budget.clean(bank.get(key))
        if value:
            label_key = "bank" if key == "bank_name" else key
            bank_values.append((budget.clean(bank_labels.get(label_key)) or label_key, value))
    if bank_values:
        heading = document.add_paragraph()
        heading.paragraph_format.space_before = Pt(8)
        _add_text(
            heading,
            budget.clean(labels.get("bank_info")) or _("Bank Information"),
            font=style["heading_font"],
            size=10,
            color=style["primary"],
            locale=style["locale"],
            rtl=style["rtl"],
            bold=True,
            alignment="left",
        )
        table = document.add_table(rows=0, cols=2)
        table.style = "Table Grid"
        _set_table_direction(table, style["rtl"])
        for label, value in bank_values:
            row = table.add_row()
            for index, text in enumerate((label, value)):
                cell = row.cells[index]
                _set_cell_margins(cell)
                if index == 0:
                    _set_cell_shading(cell, style["secondary"])
                paragraph = _remove_cell_paragraph(cell)
                _add_text(
                    paragraph,
                    text,
                    font=style["body_font"],
                    size=8.5,
                    color=style["text"],
                    locale=style["locale"],
                    rtl=style["rtl"],
                    bold=index == 0,
                    alignment="left",
                )


def _validate_editable_archive(content, expected_title, expected_number):
    if not isinstance(content, (bytes, bytearray)) or len(content) > MAX_EDITABLE_DOCX_BYTES:
        raise EditableDocxError(_("The editable Word output exceeds the 100 MB limit."))
    try:
        with ZipFile(BytesIO(content)) as archive:
            if archive.testzip() is not None:
                raise ValueError("corrupt OOXML member")
            names = archive.namelist()
            if len(names) > 512:
                raise ValueError("too many OOXML members")
            if len(names) != len(set(names)):
                raise ValueError("duplicate OOXML members")
            if any(
                name.startswith(("/", "\\"))
                or "\\" in name
                or ".." in name.split("/")
                for name in names
            ):
                raise ValueError("unsafe OOXML member path")
            if sum(item.file_size for item in archive.infolist()) > MAX_EDITABLE_DOCX_BYTES * 2:
                raise ValueError("expanded OOXML is too large")
            required = {
                "[Content_Types].xml",
                "word/document.xml",
                "word/styles.xml",
                "word/_rels/document.xml.rels",
                "docProps/core.xml",
            }
            if not required.issubset(names):
                raise ValueError("missing OOXML member")
            media_members = [
                item for item in archive.infolist() if item.filename.startswith("word/media/")
            ]
            if len(media_members) > 1 or any(item.file_size > MAX_LOGO_BYTES for item in media_members):
                raise ValueError("unexpected OOXML media payload")
            parser = etree.XMLParser(
                resolve_entities=False,
                no_network=True,
                load_dtd=False,
                huge_tree=False,
            )
            roots = {}
            for name in names:
                if name.endswith(".xml") or name.endswith(".rels"):
                    roots[name] = etree.fromstring(archive.read(name), parser=parser)
            document_xml = archive.read("word/document.xml")
            core_xml = archive.read("docProps/core.xml")
            media_count = len(media_members)
    except (BadZipFile, KeyError, ValueError, etree.XMLSyntaxError) as exc:
        raise EditableDocxError(
            _("The editable Word document was created, but the DOCX archive could not be validated.")
        ) from exc

    word_ns = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    package_rel_ns = "http://schemas.openxmlformats.org/package/2006/relationships"
    drawing_ns = (
        "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
    )
    root = roots["word/document.xml"]
    page_size = root.find(".//{%s}pgSz" % word_ns)
    text = "".join(node.text or "" for node in root.iter("{%s}t" % word_ns))
    relationships_are_internal = all(
        relation.get("TargetMode") != "External"
        for rel_root in roots.values()
        for relation in rel_root.findall("{%s}Relationship" % package_rel_ns)
    )
    logo_extents = [
        extent
        for xml_root in roots.values()
        for extent in xml_root.findall(".//{%s}extent" % drawing_ns)
    ]
    logo_extents_are_safe = (
        len(logo_extents) == media_count
        and all(
            0 < int(extent.get("cx", "0")) <= int(MAX_LOGO_WIDTH_CM * _EMU_PER_CM)
            and 0 < int(extent.get("cy", "0")) <= int(MAX_LOGO_HEIGHT_CM * _EMU_PER_CM)
            for extent in logo_extents
        )
    )
    forbidden = (
        b"<w:altChunk",
        b"<w:object",
        b"<o:OLEObject",
        b"<wp:anchor",
    )
    if (
        page_size is None
        or page_size.get("{%s}w" % word_ns) != _A4_WIDTH_TWIPS
        or page_size.get("{%s}h" % word_ns) != _A4_HEIGHT_TWIPS
        or root.find(".//{%s}tbl" % word_ns) is None
        or not root.findall(".//{%s}t" % word_ns)
        or (expected_title and expected_title not in text)
        or (expected_number and expected_number not in text)
        or not relationships_are_internal
        or not logo_extents_are_safe
        or any(marker in document_xml for marker in forbidden)
        or EDITABLE_DOCX_FINGERPRINT.encode("utf-8") not in core_xml
    ):
        raise EditableDocxError(
            _("The editable Word document's text, table, or A4 structure could not be validated.")
        )


def render_editable_docx(context, language_code=None):
    """Render an editable A4 DOCX from ``_rds_document_context`` output."""
    if not isinstance(context, Mapping):
        raise EditableDocxError(_("The document data for editable Word is invalid."))
    try:
        from docx import Document
    except ImportError as exc:  # pragma: no cover - deployment dependency
        raise EditableDocxError(
            _("The python-docx package is required for editable Word output.")
        ) from exc

    budget = _TextBudget()
    theme = _mapping(context.get("theme"))
    layout_style = context.get("layout_style")
    if layout_style not in _LAYOUT_PROFILES:
        layout_style = "technology"
    resolved_code = language_code or context.get("lang_code") or "en_US"
    locale = _resolve_docx_locale(resolved_code)
    rtl = context.get("direction") == "rtl" or str(resolved_code).startswith("ar")
    primary = _color(theme.get("primary"), "17345F")
    secondary = _color(theme.get("secondary"), "F4F7FA")
    accent = _color(theme.get("accent"), "D33A35")
    profile = _LAYOUT_PROFILES[layout_style]
    style = {
        "primary": primary,
        "secondary": secondary,
        "accent": accent,
        "text": _color(theme.get("text"), "27313A"),
        "heading_font": _font(theme.get("heading_font")),
        "body_font": _font(theme.get("body_font")),
        "locale": locale,
        "rtl": rtl,
        "language_prefix": str(resolved_code).split("_", 1)[0].lower(),
        "layout_style": layout_style,
        "title_fill": {"primary": primary, "accent": accent}[profile[0]],
        "line_header_fill": {"primary": primary, "accent": accent}[profile[1]],
        "total_fill": {"primary": primary, "accent": accent}[profile[2]],
    }
    show_company = bool(context.get("show_company", True))
    logo = (
        _prepare_logo(context.get("logo_bytes"))
        if show_company
        else (b"", 0, 0)
    )

    document = Document()
    section = _configure_document(
        document,
        body_font=style["body_font"],
        heading_font=style["heading_font"],
        locale=locale,
        rtl=rtl,
        has_company_header=show_company,
    )
    _render_header(section, context, budget, style, logo)
    _render_footer(section, context, budget, style)
    _render_title(document, context, budget, style)
    _render_party_and_metadata(document, context, budget, style)
    if context.get("show_lines", True):
        _render_lines(document, context, budget, style)
    _render_totals(document, context, budget, style)
    _render_notes_and_bank(document, context, budget, style)

    expected_title = unicodedata.normalize(
        "NFC",
        _CONTROL_CHARACTERS.sub("", str(context.get("title") or _("Document"))),
    )[:MAX_TEXT_CHARS]
    expected_number = unicodedata.normalize(
        "NFC",
        _CONTROL_CHARACTERS.sub("", str(context.get("number") or "-")),
    )[:MAX_TEXT_CHARS]
    properties = document.core_properties
    properties.author = "Ranvals Software"
    properties.last_modified_by = "DocuCraft"
    properties.title = (f"{expected_title} {expected_number}").strip()[:255]
    properties.category = "Business document"

    output = BytesIO()
    try:
        document.save(output)
    except (OSError, ValueError) as exc:
        raise EditableDocxError(_("The editable Word file could not be saved.")) from exc
    content = output.getvalue()
    _validate_editable_archive(content, expected_title, expected_number)
    return content
